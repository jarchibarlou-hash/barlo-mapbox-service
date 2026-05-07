#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
BARLO Premium PPTX Generator (Push 22.2 — template 22 slides)

Charge le template Canva 'template_diagnostic_premium.pptx' (22 slides,
pre-rempli avec un cas Behalal Marcelle / 500m² Douala / 33M FCFA) et y
remplace les textes par les vraies donnees BARLO du lead courant.

Usage : python3 generate_pptx_premium.py <data.json> <template.pptx> <output.pptx>
"""
import sys
import json
import os
from pptx import Presentation


# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────

def f(flat, key, default=""):
    val = flat.get(key, default)
    if val is None or val == "":
        return str(default)
    return str(val)


def fnum(flat, key, default=0):
    try:
        v = flat.get(key, default)
        if isinstance(v, str):
            v = "".join(c for c in v if c.isdigit() or c in ".-")
        return float(v) if v not in ("", None) else default
    except (ValueError, TypeError):
        return default


def split_name(full_name):
    """Decoupe 'Valcy Vanelle' en (first='Valcy', last='Vanelle')."""
    parts = (full_name or "").strip().split()
    if not parts:
        return ("", "")
    if len(parts) == 1:
        return (parts[0], "")
    return (parts[0], " ".join(parts[1:]))


def role_label(role):
    """INTENSIFICATION → Intensification etc."""
    return {
        "INTENSIFICATION": "Intensification",
        "EQUILIBRE": "Équilibre",
        "PRUDENT": "Prudence",
    }.get(str(role).upper(), str(role).title() if role else "")


# ─────────────────────────────────────────────────────────────────────────────
# Mapping slide-par-slide (22 slides) : ancien_texte → nouveau_texte
# ─────────────────────────────────────────────────────────────────────────────

def build_slide_replacements(flat, scenarios=None):
    # ── Identifiants client ──
    client_name = f(flat, "client_name", "")
    first, last = split_name(client_name)
    city = f(flat, "city", "Douala")

    # ── Donnees site ──
    site_area = f(flat, "site_area", "500")
    budget_fcfa = f(flat, "budget_fcfa", "33M FCFA")
    program = f(flat, "program_main", "Petit collectif")

    # ── Recommandation ──
    rec = f(flat, "rec_scenario", "C")
    rec_role = f(flat, f"{rec}_role", "PRUDENT")
    rec_role_lbl = role_label(rec_role)
    rec_score = f(flat, "rec_score", "84")
    rec_units = f(flat, f"{rec}_units", "2")
    rec_sdp = f(flat, f"{rec}_sdp", "125")
    rec_cost_total = f(flat, f"{rec}_cost_total", "23M FCFA")

    # ── Per-scenario ──
    a_units = f(flat, "A_units", "3")
    a_sdp = f(flat, "A_sdp", "200")
    a_cost = f(flat, "A_cost_total", "42M FCFA")
    a_score = f(flat, "A_score", "39")
    a_role_lbl = role_label(f(flat, "A_role", "INTENSIFICATION"))
    a_m2_logt = f(flat, "A_m2_par_logt", "75")

    b_units = f(flat, "B_units", "3")
    b_sdp = f(flat, "B_sdp", "177")
    b_cost = f(flat, "B_cost_total", "33M FCFA")
    b_score = f(flat, "B_score", "71")
    b_role_lbl = role_label(f(flat, "B_role", "EQUILIBRE"))
    b_m2_logt = f(flat, "B_m2_par_logt", "64")

    c_units = f(flat, "C_units", "2")
    c_sdp = f(flat, "C_sdp", "125")
    c_cost = f(flat, "C_cost_total", "23M FCFA")
    c_score = f(flat, "C_score", "84")
    c_role_lbl = role_label(f(flat, "C_role", "PRUDENT"))
    c_m2_logt = f(flat, "C_m2_par_logt", "75")

    # ── Calculs derives (gains B vs A) ──
    a_cost_num = fnum(flat, "A_cost_total", 42)  # extrait le 42 de "42M FCFA"
    b_cost_num = fnum(flat, "B_cost_total", 33)
    c_cost_num = fnum(flat, "C_cost_total", 23)
    gain_b_vs_a = max(0, int(round(a_cost_num - b_cost_num)))
    a_sdp_num = fnum(flat, "A_sdp", 200)
    b_sdp_num = fnum(flat, "B_sdp", 177)
    diff_sdp_b_vs_a = max(0, int(round(a_sdp_num - b_sdp_num)))

    # ── Depassement budget A ──
    budget_num = fnum(flat, "budget_fcfa", 33)
    if budget_num > 0 and a_cost_num > budget_num:
        depassement_pct = int(round((a_cost_num - budget_num) / budget_num * 100))
    else:
        depassement_pct = 0
    depassement_text = f"{depassement_pct}%" if depassement_pct > 0 else "0%"

    # ── Mix logements A/B/C ──
    a_logt_count = max(0, int(fnum(flat, "A_units", 3)) - 1)  # commerce + (units-1) logements
    b_logt_count = max(0, int(fnum(flat, "B_units", 3)) - 1)
    c_logt_count = max(0, int(fnum(flat, "C_units", 2)) - 1)

    return {
        # ── SLIDE 1 : COUVERTURE ──
        1: {
            "Behalal": first or "Client",
            "Marcelle": last,
            " 500 m² à ": f" {site_area} m² à ",
            "Douala": city,
            " 33M FCFA": f" {budget_fcfa}",
        },

        # ── SLIDE 2 : MANIFESTO (statique, pas de remplacement) ──
        2: {},

        # ── SLIDE 3 : CONTEXTE DU PROJET ──
        3: {
            "Terrain de 500 m² situé à Douala, avec un budget initial de 33 millions FCFA. Le projet prévoit la construction de 3 unités dans une configuration mixte logement et activité commerciale.":
                f"Terrain de {site_area} m² situé à {city}, avec un budget initial de {budget_fcfa}. Le projet prévoit la construction de {a_units} unités dans une configuration mixte logement et activité commerciale.",
        },

        # ── SLIDE 4 : LECTURE STRATÉGIQUE DU TERRAIN ──
        4: {
            "L'analyse du terrain de 500 m² révèle une zone constructible optimisée après application des retraits réglementaires. Retrait avant : 5 mètres minimum depuis la voie publique. Retraits latéraux : 3 mètres de chaque côté pour ventilation et accès. Retrait arrière : 4 mètres pour conformité urbaine. Ces contraintes définissent l'emprise maximale exploitable pour le projet.":
                f"L'analyse du terrain de {site_area} m² révèle une zone constructible optimisée après application des retraits réglementaires. Retraits : {f(flat, 'retrait_avant', '5m')} avant, {f(flat, 'retrait_lateral', '3m')} latéraux, {f(flat, 'retrait_arriere', '3m')} arrière. Mitoyenneté sur {f(flat, 'retrait_mitoyennete', '0')} côté(s). Ces contraintes définissent l'emprise constructible : {f(flat, 'retrait_emprise_constructible', '?')}.",
        },

        # ── SLIDE 5 : CONTRAINTES INVISIBLES (statique architectural) ──
        5: {},

        # ── SLIDE 6 : SCÉNARIO A ──
        6: {
            "INTENSIFICATION": a_role_lbl.upper(),
            "Programme: 3 unités | 200 m² SDP": f"Programme: {a_units} unités | {a_sdp} m² SDP",
            "Commerce RDC + 2 logements T3 de 75 m² (circulation comprises )chacun":
                f"Programme : {a_logt_count} logement(s) de {a_m2_logt} m² + commerce RDC. Configuration sur {int(fnum(flat, 'A_levels', 3)) + 1} niveaux.",
            "Coût estimé: 42M FCFA": f"Coût estimé: {a_cost}",
            "Ce scénario maximise l'exploitation du terrain avec une densité optimale, dépassant le budget initial de 27%.":
                f"Ce scénario maximise l'exploitation du terrain avec une densité optimale" + (f", dépassant le budget initial de {depassement_text}." if depassement_pct > 0 else f", dans la cible budgétaire."),
        },

        # ── SLIDE 7 : RISQUES A — DÉPASSEMENT ──
        7: {
            "Coût estimé de 42M FCFA contre un budget initial de 33M FCFA, soit un dépassement de 27% impactant la viabilité financière.":
                f"Coût estimé de {a_cost} contre un budget initial de {budget_fcfa}" + (f", soit un dépassement de {depassement_text} impactant la viabilité financière." if depassement_pct > 0 else f", aligné sur l'enveloppe."),
            "Densité de construction maximale générant des contraintes structurelles, thermiques et réglementaires accrues sur le chantier.":
                "Densité de construction maximale générant des contraintes structurelles, thermiques et réglementaires accrues sur le chantier.",
        },

        # ── SLIDE 8 : RISQUES A — SCORE ──
        8: {
            "Score global de 39/100 - Ce scénario présente un niveau de risque significatif nécessitant une attention particulière aux facteurs critiques.":
                f"Score global de {a_score}/100 — Ce scénario présente un niveau de risque {('significatif' if int(fnum(flat, 'A_score', 39)) < 60 else 'modéré')} nécessitant une attention particulière aux facteurs critiques.",
        },

        # ── SLIDE 9 : SCÉNARIO B ──
        9: {
            "Équilibre optimal entre ambition et maîtrise budgétaire. Programme de 3 unités pour 177 m² SDP, avec un coût total de 33M FCFA aligné sur le budget initial.":
                f"{role_label(f(flat, 'B_role', 'EQUILIBRE'))} entre ambition et maîtrise budgétaire. Programme de {b_units} unités pour {b_sdp} m² SDP, avec un coût total de {b_cost} à comparer au budget initial de {budget_fcfa}.",
            "Commerce RDC + 2 logements T3 de 64 m² (circulation comprises )chacun  Score de recommandation : 71/100.":
                f"Programme : {b_logt_count} logement(s) de {b_m2_logt} m² + commerce RDC. Score de recommandation : {b_score}/100.",
        },

        # ── SLIDE 10 : GAINS B vs A + RISQUES B ──
        10: {
            "Réduction de 9M FCFA sur le coût total et diminution de23 m² SDP pour une meilleure rentabilité au m².":
                f"Réduction de {gain_b_vs_a}M FCFA sur le coût total et diminution de {diff_sdp_b_vs_a} m² SDP pour une meilleure rentabilité au m².",
            "Risques modérés : complexité technique moyenne, délais maîtrisés, budget respecté, flexibilité d'usage préservée.":
                "Risques modérés : complexité technique moyenne, délais maîtrisés, budget respecté, flexibilité d'usage préservée.",
        },

        # ── SLIDE 11 : SCORE FIABILITÉ B ──
        11: {
            "Un score équilibré qui reflète un compromis optimal entre ambition et maîtrise des risques financiers et techniques.":
                f"Score de {b_score}/100 reflétant un compromis entre ambition et maîtrise des risques financiers et techniques.",
            "Score de Fiabilité : 71/100": f"Score de Fiabilité : {b_score}/100",
        },

        # ── SLIDE 12 : SCÉNARIO C ──
        12: {
            "Le scénario Prudence propose 2 unités pour 125 m² SDP, avec un coût total de 23M FCFA, nettement sous le budget initial de 33M FCFA.":
                f"Le scénario {c_role_lbl} propose {c_units} unités pour {c_sdp} m² SDP, avec un coût total de {c_cost} à comparer au budget initial de {budget_fcfa}.",
            "Commerce RDC + 1 logements T3 de 75 m² (circulation comprises )chacun  Cette approche minimise les risques financiers tout en garantissant une rentabilité optimale. en prévoyant un phasage de construction.":
                f"Programme : {c_logt_count} logement(s) de {c_m2_logt} m² + commerce RDC. Cette approche minimise les risques financiers et permet un phasage de construction.",
        },

        # ── SLIDE 13 : PROFIL RECOMMANDÉ C ──
        13: {
            "Configuration prudente validée par l'analyse multicritères : risques financiers minimisés, faisabilité technique confirmée, délais réalistes.":
                f"Configuration {c_role_lbl.lower()} validée par l'analyse multicritères : risques financiers minimisés, faisabilité technique confirmée, délais réalistes.",
            "Indicateurs clés au vert : coût maîtrisé (23M FCFA), surface efficiente (125 m² SDP), complexité réduite et marge de sécurité préservée.":
                f"Indicateurs clés : coût {c_cost}, surface {c_sdp} m² SDP, complexité réduite et marge de sécurité préservée.",
        },

        # ── SLIDE 14 : SCORE C ──
        14: {
            "Score Global : 84/100": f"Score Global : {c_score}/100",
            "Profil de risque optimal avec une excellente maîtrise des contraintes budgétaires et techniques. Ce scénario présente le meilleur équilibre risque/rendement.":
                f"Profil de risque {('optimal' if int(fnum(flat, 'C_score', 84)) >= 80 else 'maîtrisé')} avec une bonne maîtrise des contraintes budgétaires et techniques. Score : {c_score}/100.",
        },

        # ── SLIDE 15 : COMPARATIF GLOBAL ──
        15: {
            "Ce tableau synthétise les indicateurs clés des trois scénarios : SDP, surface habitable, efficacité spatiale, coût total, score de recommandation et durée de chantier. Le scénario C (Prudence) offre le meilleur équilibre risque/performance.":
                f"Ce tableau synthétise les indicateurs clés des trois scénarios : SDP, surface habitable, efficacité spatiale, coût total, score de recommandation et durée de chantier. Le scénario {rec} ({rec_role_lbl}) est recommandé avec un score de {rec_score}/100.",
        },

        # ── SLIDE 16 : ARBITRAGE ──
        16: {
            "Comparaison stratégique des trois scénarios selon les critères clés : coût total d'investissement, surface habitable générée et score de recommandation global. Le scénario C (orange) offre le meilleur équilibre risque/rendement.":
                f"Comparaison stratégique des trois scénarios selon les critères clés : coût total, surface habitable générée et score de recommandation. Le scénario {rec} ({rec_role_lbl}) offre le meilleur équilibre risque/rendement.",
        },

        # ── SLIDE 17 : ANALYSE BUDGETAIRE C ──
        17: {
            "Coût travaux : 23M FCFA pour un budget initial de 33M FCFA.":
                f"Coût travaux : {rec_cost_total} pour un budget initial de {budget_fcfa}.",
            "Honoraires architecte : 2M-3M FCFA (10%-15% des travaux)":
                f"Honoraires architecte : {f(flat, f'{rec}_hono_bas_M', '?')}M-{f(flat, f'{rec}_hono_haut_M', '?')}M FCFA (10%-15% des travaux)",
        },

        # ── SLIDE 18 : POINTS CLÉS & CHECKLIST (statique) ──
        18: {},

        # ── SLIDE 19 : ÉTUDE DE FAISABILITÉ — PHASAGE (statique) ──
        19: {},

        # ── SLIDE 20 : PÉRIMÈTRE DES ÉTUDES (statique) ──
        20: {},

        # ── SLIDE 21 : PROJECTION FINALE ──
        21: {
            "Le Scénario C représente l'aboutissement optimal de votre projet immobilier. Cette configuration prudente maximise la rentabilité tout en minimisant les risques structurels et financiers.":
                f"Le Scénario {rec} représente l'aboutissement optimal de votre projet immobilier. Cette configuration {rec_role_lbl.lower()} équilibre rentabilité et maîtrise des risques structurels et financiers.",
            "Avec un score de recommandation de 84/100, ce choix stratégique garantit une exécution maîtrisée dans le respect de votre enveloppe budgétaire dans le cas où vous voudriez phaser votre projet progressivement et le faire construire au fur et à mesure.":
                f"Avec un score de recommandation de {rec_score}/100, ce choix stratégique garantit une exécution maîtrisée dans le respect de votre enveloppe budgétaire de {budget_fcfa}, avec possibilité de phasage progressif.",
        },

        # ── SLIDE 22 : MERCI (contact statique) ──
        22: {},
    }


# ─────────────────────────────────────────────────────────────────────────────
# Application des replacements
# ─────────────────────────────────────────────────────────────────────────────

def apply_replacements_to_slide(slide, replacements):
    applied = 0
    for old_text, new_text in replacements.items():
        if not old_text:
            continue
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            tf = shape.text_frame
            for para in tf.paragraphs:
                para_text = "".join(r.text for r in para.runs)
                if old_text in para_text:
                    if para.runs:
                        new_para_text = para_text.replace(old_text, new_text)
                        para.runs[0].text = new_para_text
                        for r in para.runs[1:]:
                            r.text = ""
                        applied += 1
                        break
    return applied


def main():
    if len(sys.argv) < 4:
        print("Usage: python3 generate_pptx_premium.py <data.json> <template.pptx> <output.pptx>", file=sys.stderr)
        sys.exit(1)

    data_file = sys.argv[1]
    template = sys.argv[2]
    output = sys.argv[3]

    if not os.path.exists(data_file):
        print(f"ERROR: data file not found: {data_file}", file=sys.stderr)
        sys.exit(2)
    if not os.path.exists(template):
        print(f"ERROR: template not found: {template}", file=sys.stderr)
        sys.exit(3)

    with open(data_file, "r", encoding="utf-8") as fp:
        data = json.load(fp)

    flat = data.get("flat", {}) or {}
    scenarios = data.get("scenarios", {})

    print(f"[premium] Loading template: {template}")
    prs = Presentation(template)
    print(f"[premium] Template loaded: {len(prs.slides)} slides")

    print(f"[premium] Building replacements for lead: {flat.get('client_name', '?')}")
    slide_replacements = build_slide_replacements(flat, scenarios)

    total_applied = 0
    for i, slide in enumerate(prs.slides):
        slide_num = i + 1
        if slide_num in slide_replacements and slide_replacements[slide_num]:
            applied = apply_replacements_to_slide(slide, slide_replacements[slide_num])
            print(f"[premium] Slide {slide_num} : {applied} replacements applied")
            total_applied += applied

    print(f"[premium] Total replacements applied : {total_applied}")
    print(f"[premium] Saving output: {output}")
    prs.save(output)
    print(f"[premium] Done. File size: {os.path.getsize(output)} bytes")


if __name__ == "__main__":
    main()
